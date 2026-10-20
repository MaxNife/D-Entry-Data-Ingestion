"""Workspace features: contracts from the UI, as-is rejected rows, the fix grid, dry runs,
tables and export, read-only SQL, and the reprocess-conflict response."""

import csv
import io
import json

import pandas as pd
import pytest
from expected import REF_CUSTOMERS
from harness import fixture, reprocess, upload, upload_and_wait
from openpyxl import load_workbook

from ingest.core import contract_store, query

pytestmark = pytest.mark.integration

SCORES = {
    "contract": "class_scores",
    "description": "test",
    "source": {"formats": ["csv"], "encoding": "utf-8", "delimiter": ","},
    "target": {"table": "class_scores", "key": "row_hash", "on_conflict": "update"},
    "columns": [
        {"name": "name", "source_name": "Name", "type": "str", "required": True},
        {
            "name": "score",
            "source_name": "Score",
            "type": "int",
            "required": True,
            "min": 0,
            "max": 100,
        },
        {"name": "grade", "source_name": "Grade", "type": "str", "required": True},
    ],
    "rules": [],
}


def save(api, spec, mode="new"):
    return api.post(
        "/contracts/save",
        data={"mode": mode, "editor": "form", "spec_json": json.dumps(spec)},
        follow_redirects=False,
    )


# ---- contracts ---------------------------------------------------------------------


def test_new_contract_is_saved_as_v1_and_keeps_zero_bounds(api, contracts_dir):
    resp = save(api, SCORES)
    assert resp.status_code == 303, resp.text
    assert resp.headers["location"] == "/contracts/class_scores?v=1&saved=1"
    text = (contracts_dir / "class_scores" / "v1.yaml").read_text(encoding="utf-8")
    assert "min: 0" in text  # 0 is a real bound, not an empty value
    assert api.get("/contracts/class_scores").status_code == 200


def test_editing_writes_the_next_version_and_never_overwrites(api, contracts_dir):
    save(api, SCORES)
    v1 = (contracts_dir / "class_scores" / "v1.yaml").read_bytes()
    edited = {
        **SCORES,
        "columns": [*SCORES["columns"][:2], {**SCORES["columns"][2], "in_set": "A, B, C"}],
    }
    resp = save(api, edited, mode="edit")
    assert resp.headers["location"] == "/contracts/class_scores?v=2&saved=1"
    assert (contracts_dir / "class_scores" / "v1.yaml").read_bytes() == v1
    assert contract_store.versions_of("class_scores") == [1, 2]


def test_new_contract_with_an_existing_name_is_refused(api):
    save(api, SCORES)
    resp = save(api, SCORES)
    assert resp.status_code == 400
    assert "already exists" in resp.text


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"target": {"table": "batches", "key": "row_hash"}}, "reserved"),
        ({"contract": "Bad Name"}, "Contract name"),
        ({"rules": [{"code": "X", "check": "is big", "message": "m"}]}, "not a comparison"),
    ],
)
def test_invalid_contracts_are_refused_with_a_reason(api, change, message):
    resp = save(api, {**SCORES, **change})
    assert resp.status_code == 400
    assert message in resp.text


def test_unused_version_can_be_deleted_but_a_used_one_cannot(api, db, contracts_dir):
    save(api, SCORES)
    upload_and_wait(api, fixture("clean_scores.csv"), "class_scores")
    save(api, SCORES, mode="edit")  # v2, never used
    resp = api.post("/contracts/class_scores/delete", data={"version": "2"}, follow_redirects=False)
    assert resp.status_code == 303
    assert contract_store.versions_of("class_scores") == [1]
    resp = api.post("/contracts/class_scores/delete", data={"version": "1"}, follow_redirects=False)
    assert "cannot+be+deleted" in resp.headers["location"]
    assert contract_store.versions_of("class_scores") == [1]


def test_inference_guesses_types_key_and_ignores_a_few_misfits():
    spec = contract_store.infer_contract(fixture("dirty_transactions.csv"), file_name="txns.csv")
    types = {c["source_name"]: c["type"] for c in spec["columns"]}
    assert types == {
        "TxnId": "str",
        "Amount": "float",
        "Type": "str",
        "CustomerId": "str",
        "TxnDate": "date",
        "Memo": "str",
    }
    assert spec["target"]["key"] == ["txn_id"]
    assert spec["contract"] == "txns"
    assert contract_store.check_contract(contract_store._clean(spec)) == []


# ---- rejected rows as-is and the fix grid ------------------------------------------------


def _dirty_batch(api):
    return upload_and_wait(api, fixture("dirty_transactions.csv"), "txn_test@1")


def test_rejected_rows_download_matches_the_original_file(api, db):
    b = _dirty_batch(api)
    resp = api.get(f"/api/batches/{b['batch_id']}/rejected-rows?format=csv")
    assert resp.status_code == 200
    got = list(csv.reader(io.StringIO(resp.text)))
    source = list(csv.reader(fixture("dirty_transactions.csv").open(encoding="utf-8")))
    assert got[0] == source[0]  # same headers, same order
    rejected_numbers = sorted({r["row_number"] for r in db.rejections(b["batch_id"])})
    assert got[1:] == [source[n] for n in rejected_numbers]  # row n is line n after the header
    xlsx = api.get(f"/api/batches/{b['batch_id']}/rejected-rows?format=xlsx")
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert [c.value for c in next(sheet.iter_rows())] == source[0]


def test_fixing_rows_creates_a_new_batch_and_leaves_the_original_alone(api, db):
    b = _dirty_batch(api)
    before = db.batch(b["batch_id"])
    grid = api.get(f"/batches/{b['batch_id']}/fix")
    assert grid.status_code == 200
    headers, rows = _grid_from_rejections(db, b["batch_id"])
    fixed = [_repair(headers, r) for r in rows]
    payload = {
        "headers": headers,
        "rows": [{"row_number": i + 1, "values": v, "include": True} for i, v in enumerate(fixed)],
    }
    check = api.post(
        f"/batches/{b['batch_id']}/fix/validate",
        data={"rows_json": json.dumps(payload), "contract": "txn_test@v2"},
    )
    assert check.status_code == 200
    summary = check.text.split('data-testid="fix-summary">')[1].split("</div>")[0]
    assert "selected rows pass" in summary, " ".join(summary.split())
    resp = api.post(
        f"/batches/{b['batch_id']}/fix/submit",
        data={"rows_json": json.dumps(payload), "contract": "txn_test@v2"},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    new_id = resp.headers["location"].split("/batches/")[1].split("?")[0]
    from harness import wait_for_batch

    new = wait_for_batch(api, new_id)
    assert new["status"] == "completed"
    assert new["accepted_rows"] == len(fixed)
    after = db.batch(b["batch_id"])
    assert (after["status"], after["accepted_rows"], after["rejected_rows"]) == (
        before["status"],
        before["accepted_rows"],
        before["rejected_rows"],
    )


def _grid_from_rejections(db, batch_id):
    originals = {}
    for r in db.rejections(batch_id):
        originals[r["row_number"]] = r["original_row"]
    headers = list(next(iter(originals.values())).keys())
    return headers, [
        [("" if originals[n].get(h) is None else str(originals[n].get(h))) for h in headers]
        for n in sorted(originals)
    ]


def _repair(headers, values):
    row = dict(zip(headers, values, strict=True))
    try:
        row["Amount"] = str(abs(float(row["Amount"])))
    except ValueError:
        row["Amount"] = "1.0"
    if row["Type"] not in ("PAYMENT", "TRANSFER", "CASH_OUT", "REFUND"):
        row["Type"] = "PAYMENT"
    if row.get("CustomerId") not in REF_CUSTOMERS:
        row["CustomerId"] = REF_CUSTOMERS[0]
    return [row[h] for h in headers]


# ---- dry runs ---------------------------------------------------------------------


def test_validate_only_loads_nothing_and_does_not_block_a_real_upload(api, db):
    with fixture("mixed.csv").open("rb") as fp:
        resp = api.post(
            "/validate",
            files={"file": ("mixed.csv", fp, "text/csv")},
            data={"contract": "txn_test@1"},
            follow_redirects=False,
        )
    assert resp.status_code == 303 and resp.headers["location"].startswith("/validations/")
    page = api.get(resp.headers["location"])  # the background check has finished by now
    assert page.status_code == 200 and 'data-status="invalid"' in page.text
    assert "Some rows fail" in page.text
    assert db.batches() == []  # no batch row
    real = upload(api, fixture("mixed.csv"), "txn_test@1")
    assert real.json()["skipped"] is False


def test_validation_history_is_filtered_by_contract(api, db):
    for name, contract in (("mixed.csv", "txn_test@1"), ("dirty_scores.csv", "score_test")):
        with fixture(name).open("rb") as fp:
            api.post(
                "/validate", files={"file": (name, fp, "text/csv")}, data={"contract": contract}
            )
    scores = api.get("/validations", params={"contract": "score_test"}).text
    assert "dirty_scores.csv" in scores and "mixed.csv" not in scores
    # the Add data page lists only the selected contract's validations
    partial = api.get("/partials/contract", params={"contract": "txn_test"}).text
    assert "mixed.csv" in partial and "dirty_scores.csv" not in partial


def test_progress_estimates_rows_before_a_run(tmp_path):
    from ingest.core.progress import estimate_rows, snapshot

    assert estimate_rows(fixture("clean_transactions.csv")) == 5
    no_newline = tmp_path / "x.csv"
    no_newline.write_text("a,b\n1,2\n3,4", encoding="utf-8")
    assert estimate_rows(no_newline) == 2
    assert snapshot(50, 100, None, False)["pct"] == 50.0
    assert snapshot(200, 100, None, False)["pct"] == 99.0  # an estimate never claims done
    assert snapshot(0, None, None, True)["pct"] == 100.0


def test_admin_system_page_reports_the_installation(api, db):
    assert api.get("/settings", follow_redirects=False).headers["location"] == "/settings/profile"
    page = api.get("/admin/system")
    assert page.status_code == 200
    assert "Reachable" in page.text and "0003_users_and_audit" in page.text


# ---- tables, export, SQL ----------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json", "parquet"])
def test_table_export_round_trips_every_row(api, db, fmt):
    upload_and_wait(api, fixture("clean_transactions.csv"), "txn_test@1")
    upload_and_wait(api, fixture("dirty_transactions.csv"), "txn_test@1")
    count = api.get("/query", params={"sql": "select count(*) from test_transactions", "run": 1})
    assert count.status_code == 200
    resp = api.get(f"/tables/test_transactions/export?format={fmt}")
    assert resp.status_code == 200
    body = io.BytesIO(resp.content)
    frame = {
        "csv": lambda: pd.read_csv(body),
        "xlsx": lambda: pd.read_excel(body),
        "json": lambda: pd.DataFrame(json.loads(resp.content)),
        "parquet": lambda: pd.read_parquet(body),
    }[fmt]()
    assert len(frame) == db.target_count("test_transactions") > 0
    assert "row_key" not in frame.columns  # system columns are opt-in


def test_tables_outside_the_catalogue_are_not_served(api, db):
    assert api.get("/tables/batches").status_code == 404
    assert api.get("/tables/batches/export").status_code == 404


def test_sql_editor_blocks_writes_and_handles_percent(api, db, db_engine):
    upload_and_wait(api, fixture("clean_transactions.csv"), "txn_test@1")
    ok = query.run_query(db_engine, "select count(*) from batches where file_name like 'clean%';")
    assert ok["rows"] == [[1]]
    for sql in (
        "delete from batches",
        "select 1; drop table batches",
        "update batches set status = 'x'",
    ):
        with pytest.raises(ValueError):
            query.run_query(db_engine, sql)
    with pytest.raises(Exception, match="read-only"):
        query.run_query(db_engine, "with d as (delete from batches returning 1) select * from d")
    assert len(db.batches()) == 1
    resp = api.post("/api/sql/run", json={"sql": "drop table batches"})
    assert resp.status_code == 400 and "Only read queries" in resp.json()["error"]
    resp = api.post("/api/sql/run", json={"sql": "select * from connections"})
    assert resp.status_code == 400
    ok = api.post("/api/sql/run", json={"sql": "select file_name from batches"}).json()
    assert ok["ok"] and ok["columns"] == ["file_name"] and ok["row_count"] == 1


# ---- reprocess conflict ---------------------------------------------------------------


def test_reprocess_onto_a_version_the_file_already_has_is_a_clear_409(api, db):
    first = upload_and_wait(api, fixture("mixed.csv"), "txn_test@1")
    twin = upload_and_wait(api, fixture("mixed.csv"), "txn_test@2")
    resp = reprocess(api, first["batch_id"])
    assert resp.status_code == 409
    body = resp.json()
    assert body["error"] == "REPROCESS_CONFLICT" and body["existing_batch_id"] == twin["batch_id"]
    assert "SQL" not in body["detail"]
    page = api.get(f"/batches/{first['batch_id']}")
    assert 'data-testid="twin-batch"' in page.text
