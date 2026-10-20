"""Progress for long-running work (batch runs, dry runs): how many rows to expect, and
how far along a run is.

Row counts are an *estimate* taken before processing: for CSV, the number of lines minus
the header (a quoted field with a line break makes it slightly high); for Excel, the
sheet's recorded dimensions. Progress is clamped so an estimate never shows above 99%
until the run has really finished.
"""

from datetime import UTC, datetime
from pathlib import Path

_cache: dict[tuple[str, int, float], int | None] = {}


def estimate_rows(path: Path) -> int | None:
    """Data rows in a CSV/Excel file, or None when it cannot be read cheaply."""
    try:
        st = path.stat()
    except OSError:
        return None
    key = (str(path), st.st_size, st.st_mtime)
    if key in _cache:
        return _cache[key]
    n: int | None
    try:
        if path.suffix.lower() in (".xlsx", ".xls"):
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                n = max((wb.worksheets[0].max_row or 1) - 1, 0)
            finally:
                wb.close()
        else:
            lines, last = 0, b"\n"
            with path.open("rb") as fp:
                while block := fp.read(1 << 20):
                    lines += block.count(b"\n")
                    last = block[-1:]
            if last != b"\n":  # no trailing newline: the last line still counts
                lines += 1
            n = max(lines - 1, 0)
    except Exception:
        n = None
    _cache[key] = n
    return n


def snapshot(
    done: int, expected: int | None, started: datetime | str | None, finished: bool
) -> dict:
    """Percent, elapsed, rate and ETA for a run that has processed ``done`` rows."""
    if isinstance(started, str):
        started = datetime.fromisoformat(started)
    if started is not None and started.tzinfo is None:
        started = started.replace(tzinfo=UTC)
    elapsed = (datetime.now(UTC) - started).total_seconds() if started else 0.0
    rate = done / elapsed if elapsed > 0.5 and done else None
    pct = None
    if finished:
        pct = 100.0
    elif expected:
        pct = min(99.0, 100.0 * done / expected)
    eta = (
        (expected - done) / rate if rate and expected and expected > done and not finished else None
    )
    return {
        "done": done,
        "expected": expected,
        "pct": pct,
        "elapsed": elapsed,
        "rate": rate,
        "eta": eta,
    }
